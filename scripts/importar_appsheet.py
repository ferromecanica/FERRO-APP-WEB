"""Importa datos del Excel de AppSheet (Taller_Mec.xlsx) a la base de Ferro.

Uso:  python scripts/importar_appsheet.py <ruta al .xlsx> repuestos   (repuestos, proveedores, markups)
      python scripts/importar_appsheet.py <ruta al .xlsx> clientes    (clientes y vehículos)
      python scripts/importar_appsheet.py <ruta al .xlsx> historial   (turnos, OTs y ventas)
      python scripts/importar_appsheet.py <ruta al .xlsx> ot 10099 10100 [--contable "<Admin Taller.xlsx>"]

El modo "ot" trae solo esas órdenes: la OT con sus renglones, su venta y, si se
le pasa el Excel de la contable, el ingreso de esa venta atado a ella. Es para
las últimas OT que quedaron cargadas en AppSheet después de la mudanza. Solo
crea el cliente y el vehículo si faltan: a los que ya están no los toca.

Se puede correr varias veces: actualiza por número (no duplica). El historial
vuelve a escribir los renglones de cada OT y venta importada, así que no se
importa después de haber tocado a mano una OT vieja.

El historial NO mueve stock: el stock que ya está cargado es el de hoy.
Hace una copia de la base antes de tocar nada.
"""
import sys
from datetime import date, datetime, time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import openpyxl  # noqa: E402
from sqlalchemy import func  # noqa: E402

from app import create_app  # noqa: E402
from app.extensions import db  # noqa: E402
from app.models import (  # noqa: E402
    Categoria, Cliente, ConfigMarkup, ConsumoOT, MovimientoStock, OrdenTrabajo, Proveedor, RegistroHoras,
    Repuesto, Subcategoria, TareaOT, Turno, Vehiculo, Venta, VentaItem,
)
from app.services.stock import repuesto_varios  # noqa: E402
from app.validaciones import cuit_valido, formatear_cuit, normalizar_patente  # noqa: E402
from app.services.backup import hacer_copia  # noqa: E402


def filas(libro, hoja):
    it = libro[hoja].iter_rows(values_only=True)
    encabezado = next(it)
    return [dict(zip(encabezado, r)) for r in it if r and r[0] is not None]


def texto(valor):
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    t = str(valor).strip()
    return t or None


def asegurar_proveedor(nombre):
    if nombre and not Proveedor.query.filter(func.lower(Proveedor.nombre) == nombre.lower()).first():
        db.session.add(Proveedor(nombre=nombre))


def importar_repuestos(libro):
    cats = {c.nombre.upper(): c for c in Categoria.query.all()}
    subs = {(s.categoria.nombre.upper(), s.nombre.upper()): s for s in Subcategoria.query.all()}
    nuevos = actualizados = 0
    for r in filas(libro, "Repuestos"):
        rid = int(r["ID_Repuesto"])
        if rid == Repuesto.ID_VARIOS:
            continue
        cat_nombre = (texto(r["Categoria"]) or "").upper()
        sub_nombre = (texto(r["Sub_categoria"]) or "").upper()
        if cat_nombre and cat_nombre not in cats:
            cats[cat_nombre] = Categoria(nombre=cat_nombre)
            db.session.add(cats[cat_nombre])
        if cat_nombre and sub_nombre and (cat_nombre, sub_nombre) not in subs:
            subs[(cat_nombre, sub_nombre)] = Subcategoria(nombre=sub_nombre, categoria=cats[cat_nombre])
            db.session.add(subs[(cat_nombre, sub_nombre)])
        proveedor = texto(r["Proveedor"])
        asegurar_proveedor(proveedor)

        descuento = r["Descuento_Oferta"] or None
        costo = float(r["Precio_Costo"] or 0)
        rep = db.session.get(Repuesto, rid)
        if rep is None:
            rep = Repuesto(id=rid, stock_actual=0)
            db.session.add(rep)
            nuevos += 1
        else:
            actualizados += 1
        rep.nombre = texto(r["Nombre"]) or f"Repuesto {rid}"
        rep.proveedor = proveedor
        rep.marca = texto(r["Marca"])
        rep.nro_parte = texto(r["Nro_Parte"])
        rep.codigo_barras = texto(r["Codigo_Barras"])
        rep.cod_proveedor = texto(r["Cod_Proveedor"])
        rep.marca_proveedor = texto(r["Marca_RSF"])
        rep.categoria = cats.get(cat_nombre)
        rep.subcategoria = subs.get((cat_nombre, sub_nombre))
        rep.costo_lista = round(costo / (1 - descuento), 2) if descuento and descuento < 1 else costo
        rep.descuento_oferta = descuento
        rep.precio_costo = costo
        rep.costo_manual = r["Costo_Manual"] is not None
        rep.markup = r["Markup"] or None
        rep.precio_venta = float(r["Precio_Venta"] or 0)  # tal cual AppSheet (no se recalcula)
        rep.comp_marca = texto(r["Comp_Auto_Marca"])
        rep.comp_modelo = texto(r["Comp_Auto_Modelo"])
        rep.comp_motor = texto(r["Comp_Auto_Motor"])
        rep.detalle = texto(r["Detalle"])

        stock = float(r["Stock_Actual"] or 0)
        diferencia = stock - (rep.stock_actual or 0)
        if diferencia:
            rep.stock_actual = stock
            db.session.flush()
            db.session.add(MovimientoStock(repuesto_id=rid, cantidad=diferencia, tipo="Ajuste",
                                           detalle="Stock importado de AppSheet"))
    return nuevos, actualizados


def importar_markups(libro):
    ConfigMarkup.query.delete()
    for m in filas(libro, "Config_Markups"):
        proveedor = texto(m["Proveedor"])
        asegurar_proveedor(proveedor)
        db.session.add(ConfigMarkup(proveedor=proveedor, marca_envase=texto(m["Marca_Envase"]), markup=m["Markup"]))
    return ConfigMarkup.query.count()


def id_cliente(codigo):
    """'CLI-003' → 3 (el número de cliente se conserva)."""
    t = texto(codigo) or ""
    return int(t.split("-")[-1]) if t.split("-")[-1].isdigit() else None


def guardar_cliente(c):
    """Escribe una fila de la hoja Clientes. Devuelve (cliente, era_nuevo)."""
    cid = id_cliente(c["ID_Cliente"])
    if cid is None:
        return None, False
    cli = db.session.get(Cliente, cid)
    nuevo = cli is None
    if nuevo:
        cli = Cliente(id=cid)
        db.session.add(cli)
    cli.nombre = texto(c["Nombre_Completo"]) or f"Cliente {cid}"
    cli.telefono = texto(c["Telefono"])
    cli.email = (texto(c["Email"]) or "").lower() or None
    cli.direccion = texto(c["Direccion"])
    cli.notas = texto(c["Notas"])
    cuit = texto(c["CUIT"])
    cli.cuit = formatear_cuit(cuit) if cuit and cuit_valido(cuit) else cuit
    cli.condicion_iva = texto(c["Condicion_IVA"]) or "Consumidor Final"
    return cli, nuevo


def guardar_vehiculo(v):
    """Escribe una fila de la hoja Vehiculos. Devuelve (vehículo, era_nuevo)."""
    patente = normalizar_patente(texto(v["ID_Patente"]))
    if not patente:
        return None, False
    veh = Vehiculo.query.filter_by(patente=patente).first()
    nuevo = veh is None
    if nuevo:
        veh = Vehiculo(patente=patente)
        db.session.add(veh)
    cid = id_cliente(v["ID_Cliente"])
    veh.cliente_id = cid if cid and db.session.get(Cliente, cid) else None
    veh.marca, veh.modelo, veh.motor = texto(v["Marca"]), texto(v["Modelo"]), texto(v["Motor"])
    veh.traccion, veh.color = texto(v["Tracción"]), texto(v["Color"])
    veh.vin, veh.ecu_marca, veh.ecu_modelo = texto(v["VIN"]), texto(v["ECU_Marca"]), texto(v["ECU_Modelo"])
    anio, km = v["Año"], v["Kilometraje"]
    veh.anio = int(anio) if isinstance(anio, (int, float)) else None
    veh.kilometraje = int(km) if isinstance(km, (int, float)) else None
    db.session.flush()
    return veh, nuevo


def importar_clientes(libro):
    nuevos = actualizados = 0
    for c in filas(libro, "Clientes"):
        cli, nuevo = guardar_cliente(c)
        if cli is None:
            continue
        nuevos, actualizados = nuevos + nuevo, actualizados + (not nuevo)
    return nuevos, actualizados


def importar_vehiculos(libro):
    nuevos = actualizados = 0
    for v in filas(libro, "Vehiculos"):
        veh, nuevo = guardar_vehiculo(v)
        if veh is None:
            continue
        nuevos, actualizados = nuevos + nuevo, actualizados + (not nuevo)
    return nuevos, actualizados


# ──────────────────────────────── Historial ─────────────────────────────────


def numero(valor):
    """El Excel de AppSheet trae los números como texto ('10051.0')."""
    if valor is None or valor == "":
        return None
    try:
        return float(str(valor).replace(",", "."))
    except ValueError:
        return None


def entero(valor):
    n = numero(valor)
    return int(n) if n is not None else None


def fecha_de(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    for formato in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(valor).strip(), formato).date()
        except (ValueError, TypeError):
            continue
    return None


def hora_de(valor):
    if isinstance(valor, datetime):
        return valor.time()
    if isinstance(valor, time):
        return valor
    try:
        return datetime.strptime(str(valor).strip(), "%H:%M:%S").time()
    except (ValueError, TypeError):
        return None


def si_no(valor):
    return str(valor).strip().lower() in ("sí", "si", "yes", "y", "true", "1")


ESTADOS_VIEJOS = {"Terminado": "Finalizada", "En Proceso": "En proceso", "Pendiente": "Pendiente"}


def importar_turnos(libro):
    """Devuelve {ID_Turno de AppSheet: turno}. Los reconoce por fecha + hora + de quién es."""
    por_codigo = {}
    for t in filas(libro, "Turnos"):
        fecha = fecha_de(t.get("Fecha"))
        if not fecha:
            continue
        cid = id_cliente(t.get("ID_Cliente"))
        cliente = db.session.get(Cliente, cid) if cid else None
        hora = hora_de(t.get("Hora"))
        turno = Turno.query.filter_by(fecha=fecha, hora=hora, cliente_id=cliente.id if cliente else None).first()
        if turno is None:
            turno = Turno(fecha=fecha, hora=hora, cliente=cliente)
            db.session.add(turno)
        patente = normalizar_patente(texto(t.get("ID_Patente")))
        turno.vehiculo = Vehiculo.query.filter_by(patente=patente).first() if patente else None
        turno.telefono = texto(t.get("Telefono"))
        turno.motivo = texto(t.get("Motivo"))
        turno.observaciones = texto(t.get("Observaciones"))
        turno.estado = texto(t.get("Estado")) or "Confirmado"
        turno.duracion_valor = numero(t.get("Duracion_Valor"))
        turno.duracion_unidad = texto(t.get("Duracion_Unidad"))
        if not turno.cliente and not turno.contacto:
            turno.contacto = "Sin nombre"
        db.session.flush()
        por_codigo[texto(t.get("ID_Turno"))] = turno
    return por_codigo


def importar_ots(libro, turnos, numeros=None):
    """Con `numeros`, solo esas órdenes; sin nada, todas las del Excel."""
    nuevas = actualizadas = 0
    ots = {}
    for o in filas(libro, "Ordenes_trabajo"):
        numero_ot = entero(o.get("ID_OT"))
        if numeros is not None and numero_ot not in numeros:
            continue
        patente = normalizar_patente(texto(o.get("ID_Patente")))
        vehiculo = Vehiculo.query.filter_by(patente=patente).first() if patente else None
        if numero_ot is None or vehiculo is None:
            continue  # sin auto no hay OT: se importa después de clientes y vehículos
        ot = db.session.get(OrdenTrabajo, numero_ot)
        if ot is None:
            ot = OrdenTrabajo(id=numero_ot, vehiculo=vehiculo)
            db.session.add(ot)
            nuevas += 1
        else:
            actualizadas += 1
        cid = id_cliente(o.get("ID_Cliente"))
        ot.vehiculo = vehiculo
        ot.cliente = db.session.get(Cliente, cid) if cid else None
        ot.fecha_ingreso = fecha_de(o.get("Fecha_Ingreso")) or date.today()
        ot.fecha_fin = fecha_de(o.get("Fecha_Fin"))
        ot.km_entrada = entero(o.get("Km_entrada"))
        ot.km_proximo_service = entero(o.get("Km_proximo_service"))
        ot.detalle = texto(o.get("Detalle_Trabajo"))
        ot.estado = ESTADOS_VIEJOS.get(texto(o.get("Estado")), "Finalizada")
        ot.presupuesto_cliente = numero(o.get("Presupuesto_Cliente"))
        # En AppSheet la OT abierta ya trae Total_Cobrado en 0: así mostraría
        # "$ 0" cobrado en vez de "—" hasta que se cierre
        ot.total_cobrado = numero(o.get("Total_Cobrado")) if ot.fecha_fin else None
        ot.clasificacion_cierre = texto(o.get("Clasificacion_Cierre"))
        ot.link_reporte = texto(o.get("Link_Reporte"))
        ot.otros = texto(o.get("Otros"))
        for campo, columna in [("aceite_motor", "Aceite_Motor"), ("aceite_caja", "Aceite_Caja"),
                               ("aceite_diferencial", "Aceite_Diferencial"), ("filtro_aceite", "Filtro_Aceite"),
                               ("filtro_aire", "Filtro_Aire"), ("filtro_habitaculo", "Filtro_habitaculo"),
                               ("filtro_combustible", "Filtro_Combustible"), ("scaneo", "Scaneo")]:
            setattr(ot, campo, si_no(o.get(columna)))
        for campo, columna in [("aceite_motor_detalle", "Aceite_Motor_Detalle"),
                               ("aceite_caja_detalle", "Aceite_Caja_Detalle"),
                               ("aceite_diferencial_detalle", "Aceite_Diferencial_Detalle")]:
            setattr(ot, campo, texto(o.get(columna)))
        turno = turnos.get(texto(o.get("ID_Turno_Origen")))
        if turno is not None and turno.ot is None:
            turno.ot = ot
            turno.estado = "Ingresado"
        db.session.flush()
        ots[numero_ot] = ot
    return ots, nuevas, actualizadas


def importar_renglones_ot(libro, ots):
    """Consumos, horas y tareas. Se borran los de las OT importadas y se vuelven a escribir."""
    repuesto_varios()  # el "Varios / Mano de Obra" (99) lleva la mano de obra de las OT viejas
    numeros = list(ots)
    for modelo in (ConsumoOT, RegistroHoras, TareaOT):
        modelo.query.filter(modelo.ot_id.in_(numeros)).delete(synchronize_session=False)

    consumos = horas = tareas = 0
    for c in filas(libro, "Consumos_OT"):
        ot = ots.get(entero(c.get("ID_OT")))
        repuesto = db.session.get(Repuesto, entero(c.get("ID_Repuesto")) or 0)
        if ot is None or repuesto is None:
            continue
        db.session.add(ConsumoOT(
            ot=ot, repuesto=repuesto,
            descripcion=texto(c.get("Detalle_Varios")) or texto(c.get("Nombre_Registro")) or repuesto.nombre,
            cantidad=numero(c.get("Cantidad")) or 0,
            precio_unitario=numero(c.get("Precio_Unitario")) or 0,
            precio_costo=numero(c.get("Precio_Costo")) or 0,
            fecha=fecha_de(c.get("Fecha_Consumo")) or ot.fecha_ingreso,
        ))
        consumos += 1

    for h in filas(libro, "Registro_Horas"):
        ot = ots.get(entero(h.get("ID_OT")))
        if ot is None:
            continue
        db.session.add(RegistroHoras(
            ot=ot, fecha=fecha_de(h.get("Fecha")) or ot.fecha_ingreso,
            mecanico=texto(h.get("Mecanico")) or "Iván", horas=numero(h.get("Horas")) or 0,
            detalle=texto(h.get("Detalle")),
        ))
        horas += 1

    for t in filas(libro, "Tareas_extra_OT"):
        ot = ots.get(entero(t.get("ID_OT")))
        descripcion = texto(t.get("Descripcion_Tarea"))
        if ot is None or not descripcion:
            continue
        db.session.add(TareaOT(ot=ot, descripcion=descripcion))
        tareas += 1
    return consumos, horas, tareas


def importar_ventas(libro, ots, solo_de_ot=False):
    """Las ventas viejas, con su renglón. Las de OT se reconocen por la OT; las de mostrador por fecha y total.

    Con `solo_de_ot` deja afuera las de mostrador: es para el modo "ot", que
    trae únicamente lo que cuelga de las órdenes pedidas.
    """
    detalles = {}
    for d in filas(libro, "Detalle_Venta"):
        detalles.setdefault(texto(d.get("ID_Venta")), []).append(d)

    nuevas = actualizadas = 0
    por_origen = {}
    for v in filas(libro, "Ventas"):
        fecha = fecha_de(v.get("Fecha"))
        if not fecha:
            continue
        ot = ots.get(entero(v.get("ID_OT")))
        if solo_de_ot and ot is None:
            continue
        cid = id_cliente(v.get("Cliente"))
        cliente = db.session.get(Cliente, cid) if cid else None
        total = numero(v.get("Total_Venta")) or 0
        if ot is not None:
            venta = Venta.query.filter_by(ot_id=ot.id).first()
        else:
            candidatas = Venta.query.filter_by(fecha=fecha, ot_id=None,
                                               cliente_id=cliente.id if cliente else None).all()
            # el total del Excel puede diferir en centavos de la suma de los renglones
            venta = candidatas[0] if len(candidatas) == 1 else next(
                (x for x in candidatas if abs(x.total - total) < 1), None)
        if venta is None:
            venta = Venta(fecha=fecha)
            db.session.add(venta)
            nuevas += 1
        else:
            actualizadas += 1
            VentaItem.query.filter_by(venta_id=venta.id).delete(synchronize_session=False)
        venta.fecha, venta.cliente, venta.ot = fecha, cliente, ot
        venta.metodo_pago = texto(v.get("Metodo_Pago"))
        venta.tipo_comprobante = texto(v.get("Tipo_Comprobante")) or "X"
        venta.link_comprobante = texto(v.get("Link_Factura"))
        db.session.flush()
        por_origen[texto(v.get("ID_Venta"))] = venta
        items = []
        for d in detalles.get(texto(v.get("ID_Venta")), []):
            repuesto_id = entero(d.get("ID_Repuesto"))
            item = VentaItem(
                venta=venta,
                repuesto_id=repuesto_id if db.session.get(Repuesto, repuesto_id or 0) else None,
                descripcion=texto(d.get("Descripcion_Manual")) or texto(d.get("Texto_Articulo")),
                cantidad=numero(d.get("Cantidad")) or 1,
                precio_unitario=numero(d.get("Precio_Unitario")) or numero(d.get("Precio_Manual")) or 0,
                costo_unitario=numero(d.get("Costo_Unitario")) or 0,
            )
            db.session.add(item)
            items.append(item)
        # En AppSheet el costo de materiales no está en el renglón sino en la venta o en la OT
        if items and not any(i.costo_unitario for i in items):
            costo = numero(v.get("Costo_Total_Venta"))
            if costo is None and ot is not None:
                costo = ot.costo_repuestos
            if costo:
                items[0].costo_unitario = costo / (items[0].cantidad or 1)
    return nuevas, actualizadas, por_origen


# ────────────────────── Una OT suelta (después de la mudanza) ───────────────────────


def asegurar_dueno(libro, numeros):
    """Crea el cliente y el vehículo de esas OT si todavía no están en la base.

    A los que ya existen no los toca: en Ferro pueden haberse corregido a mano.
    """
    filas_ot = [o for o in filas(libro, "Ordenes_trabajo") if entero(o.get("ID_OT")) in numeros]
    clientes = {texto(o.get("ID_Cliente")) for o in filas_ot}
    patentes = {normalizar_patente(texto(o.get("ID_Patente"))) for o in filas_ot}
    nuevos = []
    for c in filas(libro, "Clientes"):
        if texto(c["ID_Cliente"]) in clientes and db.session.get(Cliente, id_cliente(c["ID_Cliente"])) is None:
            cli, _ = guardar_cliente(c)
            nuevos.append(f"cliente {cli.nombre}")
    db.session.flush()
    for v in filas(libro, "Vehiculos"):
        patente = normalizar_patente(texto(v["ID_Patente"]))
        if patente in patentes and Vehiculo.query.filter_by(patente=patente).first() is None:
            guardar_vehiculo(v)
            nuevos.append(f"vehículo {patente}")
    db.session.flush()
    return nuevos


def importar_ot_suelta(libro, numeros, contable=None):
    """Las OT pedidas con todo lo que cuelga: renglones, venta e ingreso contable."""
    faltaban = asegurar_dueno(libro, numeros)
    ots, ot_n, ot_a = importar_ots(libro, {}, numeros)
    consumos, horas, tareas = importar_renglones_ot(libro, ots)
    db.session.flush()
    db.session.expire_all()  # los renglones recién escritos, para que el costo de la OT salga bien
    v_n, v_a, ventas = importar_ventas(libro, ots, solo_de_ot=True)

    m_n = m_a = 0
    if contable is not None:
        from importar_contable import importar_movimientos  # noqa: E402  (mismo directorio)
        m_n, m_a = importar_movimientos(contable, ventas=ventas)

    db.session.commit()
    if faltaban:
        print("  Faltaban en la base y se crearon: " + ", ".join(faltaban))
    print(f"✓ OTs: {ot_n} nuevas, {ot_a} actualizadas ({', '.join(str(n) for n in sorted(ots))})")
    print(f"  Renglones: {consumos} consumos, {horas} registros de horas, {tareas} tareas")
    print(f"  Ventas: {v_n} nuevas, {v_a} actualizadas · Ingresos contables: {m_n} nuevos, {m_a} actualizados")
    for n in sorted(ots):
        ot = ots[n]
        venta = ot.ventas[0] if ot.ventas else None
        plata = f"venta ${venta.total:,.0f} ({venta.metodo_pago})".replace(",", ".") if venta else "sin venta"
        quien = ot.cliente.nombre if ot.cliente else "sin cliente"
        print(f"  OT {n} · {ot.vehiculo.patente} · {quien} · {ot.estado} · {plata}")
    sin_importar = sorted(numeros - set(ots))
    if sin_importar:
        print(f"  ⚠ No estaban en el Excel (o sin patente): {sin_importar}")
    print("  (las fotos de la OT quedan en AppSheet: hay que bajarlas de Drive aparte)")


if __name__ == "__main__":
    if len(sys.argv) < 3 or sys.argv[2] not in ("repuestos", "clientes", "historial", "ot"):
        sys.exit(__doc__)
    libro = openpyxl.load_workbook(sys.argv[1], data_only=True, read_only=True)
    app = create_app()
    with app.app_context():
        print("Copia de seguridad:", hacer_copia(app.config["SQLALCHEMY_DATABASE_URI"]).name)
        if sys.argv[2] == "ot":
            resto = sys.argv[3:]
            contable = None
            if "--contable" in resto:
                i = resto.index("--contable")
                contable = openpyxl.load_workbook(resto[i + 1], data_only=True, read_only=True)
                resto = resto[:i] + resto[i + 2:]
            numeros = {int(n) for n in resto if n.isdigit()}
            if not numeros:
                sys.exit("Decime qué OT importar, por número:  ... ot 10099 10100")
            importar_ot_suelta(libro, numeros, contable)
            sys.exit(0)
        if sys.argv[2] == "historial":
            turnos = importar_turnos(libro)
            ots, ot_n, ot_a = importar_ots(libro, turnos)
            consumos, horas, tareas = importar_renglones_ot(libro, ots)
            db.session.flush()
            db.session.expire_all()  # los renglones recién escritos, para que el costo de la OT salga bien
            v_n, v_a, _ = importar_ventas(libro, ots)
            db.session.commit()
            print(f"✓ Turnos: {len(turnos)} · OTs: {ot_n} nuevas, {ot_a} actualizadas")
            print(f"  Renglones: {consumos} consumos, {horas} registros de horas, {tareas} tareas")
            print(f"  Ventas: {v_n} nuevas, {v_a} actualizadas")
            print(f"  Total en la base: {OrdenTrabajo.query.count()} OTs, {Venta.query.count()} ventas, "
                  f"{Turno.query.count()} turnos")
            print("  (las fotos de las OT quedan en AppSheet: hay que bajarlas de Drive aparte)")
            sys.exit(0)
        if sys.argv[2] == "clientes":
            c_n, c_a = importar_clientes(libro)
            db.session.flush()
            v_n, v_a = importar_vehiculos(libro)
            db.session.commit()
            sin_dueno = Vehiculo.query.filter(Vehiculo.cliente_id.is_(None)).count()
            print(f"✓ Clientes: {c_n} nuevos, {c_a} actualizados · Vehículos: {v_n} nuevos, {v_a} actualizados"
                  f" ({sin_dueno} sin dueño)")
            sys.exit(0)
        nuevos, actualizados = importar_repuestos(libro)
        reglas = importar_markups(libro)
        db.session.commit()
        print(f"✓ Repuestos: {nuevos} nuevos, {actualizados} actualizados · reglas de markup: {reglas}")
        print(f"  Total en la base: {Repuesto.query.filter(Repuesto.id != Repuesto.ID_VARIOS).count()} repuestos, "
              f"{Proveedor.query.count()} proveedores")
